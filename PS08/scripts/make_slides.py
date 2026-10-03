"""Build slides/PS08.pptx from src/ and results/.

Every number on a slide or in the speaker notes is read from results/ here,
at build time; every listing is cut from src/ or results/.  Nothing measured
is typed by hand.  Run scripts/run_all.sh first (it checks the claims the
text below makes), then scripts/build_slides.ps1, which calls this.
"""
import math
import re
import textwrap
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from results import LOCKS, THREADS, RESULTS, bench, blocks, env, median, parse, runs

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
TEMPLATE = ROOT / "UTF-8_PS8_2025.pptx"
OUT = ROOT / "slides" / "PS08.pptx"

MEMBERS = [  # (id, name, the parts they present), in speaking order
    ("67070501075", "Siriwan Yindeephot", ["How We Evaluate", "Conclusion"]),
    ("67070501018", "Tithinan Sobking", ["Peterson's Algorithm"]),
    ("67070501005", "Kiatisak Markmeeshap", ["Load-Linked / Store-Conditional"]),
    ("67070501021", "Thanaboon Tikaew", ["MCS Lock"]),
    ("67070501040", "Worawut Sereethai", ["Comparison: Throughput"]),
    ("67070501059", "Chanon Lhumsa-ard", ["Comparison: Summary"]),
]
PRESENTER = {part: (sid, name) for sid, name, parts in MEMBERS for part in parts}
LABEL = {"peterson": "Peterson", "llsc": "CAS spin lock (x86 proxy for LL/SC)", "mcs": "MCS",
         "tas": "TAS (lecture)", "mutex": "pthread_mutex (lecture)"}
ACCENT = RGBColor(0xC0, 0x39, 0x2B)    # the red the template uses for remarks
CODE_FONT = "Consolas"

B = bench()
E = env()
NCPU = E["nproc"]


# ------------------------------------------------------------- numbers ---
def tp(lock, t, mode="all-cpus"):
    return median(B[(lock, t, mode)], "throughput")


def fair(lock, t, mode="all-cpus"):
    return median(B[(lock, t, mode)], "fairness")


def rate(x):
    """Acquisitions per second, short: '17 M', '4.1 M', '89 k'."""
    if x >= 10e6:
        return f"{x / 1e6:.0f} M"
    if x >= 1e6:
        return f"{x / 1e6:.1f} M"
    return f"{x / 1e3:.0f} k"


def first_run(lock, t, mode):
    """Run 1 of a configuration, for quoting one concrete set of counts."""
    return B[(lock, t, mode)][0]


pet = runs("1_peterson.txt", r".")
nofence = runs("1_peterson.txt", r"peterson_nofence all-cpus")
nofence1 = runs("1_peterson.txt", r"peterson_nofence one-cpu")
fenced = runs("1_peterson.txt", r"peterson_fence all-cpus")
atomic = runs("1_peterson.txt", r"^peterson all-cpus")
arm = runs("2_llsc.txt", r".")
# every run of a build that is meant to be correct (the no-fence demo is not)
good_runs = [r for rs in B.values() for r in rs] + fenced + atomic + arm
N_GOOD = len(good_runs)

SECS = re.search(r"\$ \S+ \d+ (\S+)", (RESULTS / "bench.txt").read_text()).group(1)
REPS = len(B[("mutex", 1, "all-cpus")])


# ------------------------------------------------------------- sources ---
def cut(path, start, end=r"^\}"):
    """Lines of `path` from the first matching `start` through `end`."""
    lines = Path(path).read_text().splitlines()
    i = next(n for n, l in enumerate(lines) if re.search(start, l))
    j = next(n for n in range(i, len(lines)) if re.search(end, lines[n]))
    return "\n".join(lines[i:j + 1])


def disasm(name):
    """objdump transcript without the command line and the address column."""
    lines = (RESULTS / name).read_text().splitlines()[1:]
    out = []
    for l in lines:
        if "\t" in l:
            l = l.split("\t", 1)[1]
        # drop '// #1' and x86 '# 4270 <turn>' comments; keep ARM '#0xb40'
        # and objdump's nearest-symbol guesses such as <_nl_global_locale+0x50>
        l = re.sub(r"\s+<(?!lock_acquire)[^>]*>", "", l)
        out.append(re.sub(r"\s+//.*|\s+# .*", "", l).replace("\t", " "))
    return "\n".join(out)


# --------------------------------------------------------- slide tools ---
prs = Presentation(TEMPLATE)
sld_ids = prs.slides._sldIdLst
for sid in list(sld_ids):                      # drop the template's examples
    prs.part.drop_rel(sid.rId)
    sld_ids.remove(sid)

L_TITLE, L_CONTENT, L_SECTION, L_TITLE_ONLY = (prs.slide_layouts[i] for i in (0, 1, 2, 5))


def new_slide(title, notes, layout=None):
    s = prs.slides.add_slide(layout or L_TITLE_ONLY)
    s.shapes.title.text = title
    s.notes_slide.notes_text_frame.text = notes
    n = len(prs.slides)
    box = s.shapes.add_textbox(Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.text, p.alignment = str(n), PP_ALIGN.RIGHT
    p.runs[0].font.size = Pt(12)
    p.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return s


def text(slide, x, y, w, h, items, size=20, font=None, bullet=True):
    """A text box; items are strings, or (string, level), or (string, level, colour)."""
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for k, item in enumerate(items):
        s, lvl, colour = (item + (0, None))[:3] if isinstance(item, tuple) else (item, 0, None)
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        mark = ("• " if lvl == 0 else "– ") if bullet and s else ""
        p.text = mark + s
        p.level = lvl
        p.space_after = Pt(4 if font else 8)
        for r in p.runs:
            r.font.size = Pt(size - 2 * lvl)
            if font:
                r.font.name = font
            if colour:
                r.font.color.rgb = colour
    return tf


def code(slide, x, y, w, h, src, size=16):
    """A grey listing box; the font shrinks until the longest line and all
    lines fit, so a listing never wraps or spills out of its box."""
    lines = textwrap.dedent(src).split("\n")
    fit_w = (w - 0.25) * 72 / (0.56 * max(map(len, lines)))   # Consolas: 0.55 em wide
    fit_h = (h - 0.15) * 72 / (1.2 * len(lines))
    size = math.floor(2 * min(size, fit_w, fit_h)) / 2
    tf = text(slide, x, y, w, h, lines, size=size, font=CODE_FONT, bullet=False)
    tf.word_wrap = False
    for p in tf.paragraphs:
        p.space_after = Pt(0)
    box = slide.shapes[-1]
    box.height = Inches(min(h, len(lines) * 1.2 * size / 72 + 0.2))   # hug the listing
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF4, 0xF4, 0xF4)
    return tf


def table(slide, x, y, w, rows, widths, size=14, row_h=0.4):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(w), Inches(row_h * len(rows)))
    t = shape.table
    for c, cw in enumerate(widths):
        t.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
    return t


def log_axis(chart, values):
    """Log value axis spanning whole decades around the data, with the
    category axis along the bottom rather than crossing at 1."""
    scaling = chart.value_axis._element.find(qn("c:scaling"))
    lb = OxmlElement("c:logBase")
    lb.set("val", "10")
    scaling.insert(0, lb)
    va = chart.value_axis
    va.minimum_scale = 10 ** math.floor(math.log10(min(values)))
    va.maximum_scale = 10 ** math.ceil(math.log10(max(values)))
    va.crosses = XL_AXIS_CROSSES.MINIMUM


def chart(slide, x, y, w, h, kind, categories, series, cat_title, size=12, labels=False,
          value_title="million acquisitions / s (log)", log=True):
    """series: [(name, [values or None])]; a log value axis unless log=False,
    which gives a 0-1 axis (for fairness)."""
    data = CategoryChartData()
    data.categories = categories
    for name, vals in series:
        data.add_series(name, vals)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    c = gf.chart
    c.font.size = Pt(size)
    c.has_legend = True
    c.legend.position = XL_LEGEND_POSITION.BOTTOM
    c.legend.include_in_layout = False
    va = c.value_axis
    if log:
        log_axis(c, [v for _, vs in series for v in vs if v is not None])
    else:
        va.minimum_scale, va.maximum_scale = 0, 1
    va.has_title = True
    va.axis_title.text_frame.text = value_title
    va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    va.tick_labels.number_format = "General"
    va.tick_labels.number_format_is_linked = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    ca = c.category_axis
    ca.has_title = True
    ca.axis_title.text_frame.text = cat_title
    ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    if labels:
        pl = c.plots[0]
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.number_format = labels if isinstance(labels, str) else "[<1]0.00;0.0"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(size)
    return c


def mps(lock, t, mode):
    """Median throughput in M acq/s, None where the lock cannot run t threads."""
    return tp(lock, t, mode) / 1e6 if (lock, t, mode) in B else None


def section(part, notes):
    """Section header naming the member who presents `part`."""
    s = new_slide(part, notes, L_SECTION)
    s.placeholders[1].text = "\n".join(PRESENTER[part])


# ============================================================== slides ===
# ---- 1 title
s = new_slide("Problem Session 08: Lock Mechanisms",
              "We are group OS InW. This deck studies three locks that were not "
              "covered in Lecture 8: Peterson's algorithm, Load-Linked/Store-Conditional, "
              "and the MCS queue lock. Each lock is judged on correctness and performance "
              "with the metrics from the lecture.", L_TITLE)
s.placeholders[1].text = ("Peterson's Algorithm · LL/SC · MCS Lock\n"
                          "CPE 333 Operating Systems, 1/2026 — Group OS InW")

# ---- 2 group
s = new_slide("Group: OS InW",
              "There are six of us. Three members each present one lock and two compare "
              "them; one member explains how we measured and draws the conclusion.",
              L_TITLE_ONLY)
for k, (sid, name, parts) in enumerate(MEMBERS):
    tf = text(s, 0.5 + 4.15 * (k % 3), 1.9 + 2.5 * (k // 3), 4.0, 2.2,
              [sid, name, ("", 0)] + [(p, 0, ACCENT) for p in parts], size=20, bullet=False)
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(0)

# ---- 3 method
section("How We Evaluate",
        "Before the locks themselves: the metrics we use and how we measured them.")

s = new_slide("Metrics from Lecture 8",
              "Lecture 8 judges a lock on four things: mutual exclusion, absence of deadlock, "
              "fairness, and performance. We turn each one into something a program can "
              "measure, so every lock is judged the same way. A measurement can only show a "
              "violation; it cannot prove that none exists, and the next slide separates the two.")
table(s, 0.6, 1.7, 12.1, [
    ["Metric (Lecture 8)", "Question", "What we measure"],
    ["Mutual exclusion", "Is only one thread inside at a time?",
     "Shared counter must equal total acquisitions: lost = total − counter, expected 0"],
    ["Absence of deadlock", "Does some waiting thread always get in?",
     "Every run must finish when the time is up (a hang would show)"],
    ["Fairness", "Does every thread get a fair chance?",
     "Fewest ÷ most acquisitions per thread (1.0 = equal, near 0 = severe imbalance)"],
    ["Performance", "How much time does the lock cost?",
     "Acquisitions per second over the measurement window, 1 to 8 threads"],
], [2.6, 4.0, 5.5], size=16, row_h=0.9)

s = new_slide("Argument vs. Experiment",
              f"A test and a proof answer different questions. The argument says why a lock "
              f"should work, under stated assumptions. The experiment only reports what we "
              f"saw: in {N_GOOD} runs of the correct builds we observed no lost update and no "
              f"hang. That is evidence, not proof, because a rare interleaving could be missed.")
table(s, 0.6, 1.7, 12.1, [
    ["Property", "Argument (why it should hold)", "Experiment (what we observed)"],
    ["Mutual exclusion", "An invariant: at most one thread is past lock_acquire",
     f"No violation observed in {N_GOOD} runs: lost = 0"],
    ["Absence of deadlock", "Some waiting thread can always proceed",
     f"No violation observed in {N_GOOD} runs: every run finished"],
    ["Fairness", "Bounded waiting or FIFO order, if the algorithm has it",
     "Acquisition counts per thread in a 1 s window; the scheduler affects them"],
], [2.6, 4.9, 4.6], size=16, row_h=0.8)
text(s, 0.6, 5.2, 12.1, 1.8, [
    "The arguments assume: sequentially consistent memory (or the right fences), a scheduler "
    "that eventually runs every thread, and a holder that calls unlock",
    "The no-fence Peterson demo is the one build that is expected to fail",
], size=18)

s = new_slide("The Benchmark Loop",
              f"Every lock runs this same loop. All threads start together at a barrier and "
              f"loop until the main thread raises stop after {SECS} second. The counter is a "
              f"plain variable, not an atomic one, so if two threads are ever inside together, "
              f"one of their updates is lost and the counter ends short.")
text(s, 0.6, 1.7, 5.4, 5.4, [
    "All threads start together at a barrier",
    f"Loop until main raises stop after {SECS} s",
    "counter is a plain variable: two threads inside together lose an update",
    "Each thread counts its own acquisitions (n) and timestamps its start and stop",
    "Window = first thread's start to last thread's stop; creating and joining threads is outside it",
    "Same loop for every lock: only lock_acquire / lock_release change",
], size=19)
code(s, 6.3, 1.7, 6.6, 5.0, cut(SRC / "bench.c", r"^static void \*worker"))

ex_h, ex_b = next((h, b) for h, b in blocks("bench.txt") if h == "tas threads=4 all-cpus run 1")
ex = parse(ex_b)
s = new_slide("Reading One Run",
              f"This is one real run: the test-and-set lock from the lecture with four threads. "
              f"Each thread reports how many times it got the lock. The total matches the "
              f"counter, so nothing was lost. Fairness is the smallest count divided by the "
              f"largest, and performance is the total divided by the time.")
code(s, 0.6, 1.7, 6.6, 3.2, f"--- {ex_h} ---\n" + ex_b.strip())
text(s, 7.5, 1.7, 5.4, 5.4, [
    f"Mutual exclusion: total {ex['total']:,} = counter {ex['counter']:,} → lost = {ex['lost']}",
    "Absence of deadlock: the run finished",
    f"Fairness: {min(ex['counts']):,} ÷ {max(ex['counts']):,} = {ex['fairness']:.3f}",
    f"Performance: {rate(ex['throughput'])} acquisitions per second",
], size=18)

s = new_slide("Experiment Setup",
              f"Each lock runs with 1, 2, 4 and 8 threads, first on all {NCPU} CPUs and then "
              f"pinned to one CPU with taskset, to see what spinning costs when the lock holder "
              f"cannot run. Every setting runs {REPS} times and the slides show the median. Two "
              f"locks from the lecture, test-and-set and pthread_mutex, run as references.")
table(s, 0.6, 1.7, 7.2, [
    ["Setting", "Values"],
    ["Locks", "Peterson, LL/SC, MCS"],
    ["References (Lecture 8)", "TAS spin lock, pthread_mutex"],
    ["Threads", "1, 2, 4, 8 (Peterson: 1, 2)"],
    ["CPUs", f"all {NCPU}, and one (taskset -c 0)"],
    ["Runs", f"{REPS} × {SECS} s per setting, median shown (spread: Run-to-Run slide)"],
], [3.0, 4.2], size=16, row_h=0.6)
text(s, 8.2, 1.7, 4.7, 5.4, [
    f"WSL2 Ubuntu, kernel {E['kernel']}",
    f"{NCPU} CPUs",
    f"gcc {E['gcc']} -O2",
    f"LL/SC on AArch64: qemu-aarch64 {E['qemu']}",
    "WSL2 vCPUs are scheduled by Windows too, so absolute numbers will differ on other machines",
    f"Tiny critical section, no work outside it, 1 s runs, {REPS} repetitions: compare locks "
    "with each other, not with other machines",
    ("One CPU shows what spinning costs when the lock holder cannot run", 0, ACCENT),
], size=16)

# ================================================================ Peterson
barrier = next(l.split("\t", 1)[1].strip() for l in
               (RESULTS / "1_peterson_disasm.txt").read_text().splitlines()
               if re.search(r"mfence|lock or", l))

section("Peterson's Algorithm",
        "Part one: Peterson's algorithm, a lock built from ordinary loads and stores.")

s = new_slide("Peterson's Algorithm: Idea",
              "Peterson's algorithm needs no special instruction. Each thread raises its flag "
              "to say it wants in, then politely gives the turn to the other thread. The thread "
              "that wrote turn last is the one that waits.")
text(s, 0.6, 1.7, 4.8, 5.4, [
    "Pure software lock for two threads: only loads and stores",
    "flag[i] = 1 — “thread i wants to enter”",
    "turn = other — “you go first”",
    "Wait while the other wants in and it is the other's turn",
    "The thread that wrote turn last waits",
    "Reference build: _Atomic ints, seq_cst; FENCE() is empty there (the seq_cst stores order the loads)",
], size=20)
code(s, 5.6, 1.6, 7.3, 5.5,
     cut(SRC / "1_peterson.c", r"^typedef _Atomic int shared_t;", r"^typedef _Atomic int shared_t;")
     + "\n" + cut(SRC / "1_peterson.c", r"^static shared_t flag", r"^static shared_t turn")
     + "\n\n" + cut(SRC / "1_peterson.c", r"^void lock_acquire")
     + "\n\n" + cut(SRC / "1_peterson.c", r"^void lock_release"))

lost_n = [r["lost"] for r in nofence]
s = new_slide("Peterson: Mutual Exclusion and Deadlock",
              f"The argument assumes sequential consistency. Suppose both threads were inside, "
              f"and take the one that wrote turn last. The other thread's flag was set before "
              f"that thread's own write of turn, so the last writer reads flag[other] == 1 and "
              f"turn == other, and it must wait. On a real x86 CPU the textbook volatile code "
              f"breaks this assumption, because a load may overtake an earlier store. Without "
              f"the fence we lost updates in {sum(l > 0 for l in lost_n)} of {len(nofence)} "
              f"runs. With the fence, or with _Atomic, we observed no violation in "
              f"{len(fenced)} runs each.")
text(s, 0.7, 1.55, 12.0, 2.5, [
    "Mutual exclusion (assuming sequential consistency): if both were inside, take the thread "
    "P that wrote turn last; the other thread's flag[] = 1 came before its own turn write, "
    "so P reads flag[other] == 1 and turn == other, and waits — contradiction",
    "No deadlock: both waiting would need turn == 1 and turn == 0 at once; and the holder must call unlock",
    ("x86 lets a load pass an earlier store (store buffer): both can read flag[other] == 0 and both enter. "
     "That breaks the assumption, not the logic", 0, ACCENT),
], size=17)
table(s, 0.7, 4.2, 12.0, [
    ["Build (2 threads)", "CPUs", "Runs with lost updates", "Lost updates per run"],
    ["volatile, no fence", f"all {NCPU}", f"{sum(l > 0 for l in lost_n)} / {len(nofence)}",
     f"{min(lost_n):,} – {max(lost_n):,}"],
    ["volatile, no fence", "one", f"{sum(r['lost'] > 0 for r in nofence1)} / {len(nofence1)}",
     f"{max(r['lost'] for r in nofence1):,}"],
    ["volatile + FENCE()", f"all {NCPU}", f"{sum(r['lost'] > 0 for r in fenced)} / {len(fenced)}",
     f"{max(r['lost'] for r in fenced):,}"],
    ["_Atomic, seq_cst (reference)", f"all {NCPU}", f"{sum(r['lost'] > 0 for r in atomic)} / {len(atomic)}",
     f"{max(r['lost'] for r in atomic):,}"],
], [3.8, 1.5, 3.2, 3.5], size=15, row_h=0.45)
text(s, 0.7, 6.55, 12.0, 0.7, [
    f"FENCE() = {barrier}, a full barrier. The volatile builds are data races in C11: "
    f"they show GCC on x86 only",
], size=14)

s = new_slide("Peterson: Performance",
              f"With one thread there is no contention and the lock is cheap. With two threads "
              f"on separate CPUs, flag and turn bounce between the cores' caches. Pinned to one "
              f"CPU it collapses: a waiting thread spins away its whole time slice, as the "
              f"lecture warned.")
chart(s, 0.6, 1.6, 6.6, 5.4, XL_CHART_TYPE.COLUMN_CLUSTERED, ["1 thread", "2 threads"],
      [(f"all {NCPU} CPUs", [mps("peterson", t, "all-cpus") for t in (1, 2)]),
       ("one CPU", [mps("peterson", t, "one-cpu") for t in (1, 2)])],
      "threads", labels=True)
text(s, 7.5, 1.8, 5.4, 5.0, [
    f"1 thread: {rate(tp('peterson', 1))} acq/s — no contention",
    f"2 threads, all CPUs: {rate(tp('peterson', 2))} acq/s — flag and turn move between cores",
    (f"2 threads, one CPU: {rate(tp('peterson', 2, 'one-cpu'))} acq/s — the waiter spins "
     f"until the scheduler switches", 0, ACCENT),
    "Busy-waiting costs a whole time slice when the other thread is not running",
], size=19)

r1 = first_run("peterson", 2, "one-cpu")
s = new_slide("Peterson: Fairness and Limitations",
              f"With both threads on their own CPU, turn hands priority to the other thread, so "
              f"neither waits for more than one entry of the other; the counts were nearly equal. "
              f"On one CPU the counts were lopsided because who gets in follows the scheduler. "
              f"The algorithm also only works for two threads and needs sequential consistency, "
              f"which means a fence on x86.")
text(s, 0.9, 1.7, 11.8, 5.4, [
    f"All CPUs, 2 threads: fairness {fair('peterson', 2):.2f} — turn gives way to the other thread, "
    f"so a waiter enters after at most one entry by the other (bounded waiting)",
    f"One CPU, 2 threads: fairness {fair('peterson', 2, 'one-cpu'):.3f} "
    f"(run 1: {min(r1['counts']):,} vs {max(r1['counts']):,}) — the imbalance comes from the "
    f"scheduler's time slices; bounded waiting is a property of the algorithm, not of this window",
    "Limitations:",
    ("Two threads only (N threads need a tournament of Peterson locks or the filter lock)", 1),
    ("Needs sequential consistency: a full fence on x86 — the textbook volatile version is broken there", 1),
    ("Spins instead of sleeping", 1),
], size=20)

# =================================================================== LL/SC
section("Load-Linked / Store-Conditional",
        "Part two: Load-Linked and Store-Conditional, a pair of hardware instructions.")

s = new_slide("LL/SC: Idea",
              "Load-Linked reads a value and starts watching the address. Store-Conditional "
              "writes only if nobody wrote that address in between. The lock spins on "
              "Load-Linked until the flag is free, then tries to claim it; if the store fails, "
              "someone else won and the thread starts over.")
text(s, 0.6, 1.7, 4.8, 5.4, [
    "LoadLinked(p): load *p and watch the address",
    "StoreConditional(p, v): store only if nobody stored to p since the LoadLinked",
    "Lock: spin on LL until 0, then SC 1; if SC fails, start over",
    "MIPS, Alpha, PowerPC, ARM (AArch64: LDXR/STXR)",
    ("x86 has no LL/SC — it offers CAS (lock cmpxchg)", 0, ACCENT),
], size=19)
code(s, 5.6, 1.6, 7.3, 5.5,
     cut(SRC / "2_llsc.c", r"^static inline int LoadLinked\(volatile int \*p\)$", r"^\}")
     + "\n\n" + cut(SRC / "2_llsc.c", r"^static inline int StoreConditional\(volatile int \*p, int v\)$", r"^\}")
     + "\n\n" + cut(SRC / "2_llsc.c", r"^void lock_acquire"))

s = new_slide("LL/SC: Mutual Exclusion and Deadlock",
              f"Two threads may both load-link a zero, but only the first store-conditional "
              f"succeeds; the second sees that the address changed and fails. We ran the AArch64 "
              f"build under qemu with {arm[0]['counts'].__len__()} threads: no lost update was observed "
              f"in {sum(r['lost'] == 0 for r in arm)} of {len(arm)} runs, which checks the "
              f"implementation and proves nothing about every interleaving. The disassembly "
              f"shows the real ldaxr and stxr instructions.")
text(s, 0.9, 1.6, 6.3, 5.5, [
    "Mutual exclusion: two threads may both LL a 0, but the first SC writes the flag, "
    "so the second SC fails",
    "No deadlock: an SC fails only if another store got in (that thread now holds the lock) "
    "or spuriously (e.g. an interrupt) — the loop retries",
    f"AArch64 build under qemu-aarch64 {E['qemu']}, {len(arm[0]['counts'])} threads: "
    f"no violation observed in {sum(r['lost'] == 0 for r in arm)} / {len(arm)} runs "
    f"(lost = 0, {min(r['total'] for r in arm):,}+ acquisitions each)",
], size=19)
code(s, 7.5, 1.7, 5.4, 4.0, disasm("2_llsc_disasm.txt"))
text(s, 7.5, (s.shapes[-1].top + s.shapes[-1].height) / 914400 + 0.15, 5.4, 1.0, [("ldaxr = Load-Linked, stxr = Store-Conditional "
                              "(writes 0 to w0 on success)", 0, ACCENT)], size=15, bullet=False)

s = new_slide("LL/SC: Performance",
              f"qemu emulates the CPU, so its speed means nothing; we time the x86 build, which "
              f"runs the same loop with compare-and-swap standing in for the LL/SC pair. These "
              f"numbers belong to a CAS spin lock on x86; they compare implementations on x86 "
              f"and say nothing about the speed of real LL/SC hardware. It behaves like the "
              f"test-and-set lock from the lecture: fast alone, slower as threads fight over "
              f"one cache line.")
chart(s, 0.6, 1.6, 7.0, 5.4, XL_CHART_TYPE.LINE_MARKERS, [str(t) for t in THREADS],
      [(f"CAS proxy, all {NCPU} CPUs", [mps("llsc", t, "all-cpus") for t in THREADS]),
       ("CAS proxy, one CPU", [mps("llsc", t, "one-cpu") for t in THREADS]),
       (f"TAS, all {NCPU} CPUs", [mps("tas", t, "all-cpus") for t in THREADS]),
       ("TAS, one CPU", [mps("tas", t, "one-cpu") for t in THREADS])],
      "threads")
text(s, 7.9, 1.7, 5.0, 5.3, [
    ("Timed on x86: a CAS spin lock, NOT hardware LL/SC (qemu speed is meaningless)", 0, ACCENT),
    f"All CPUs: {rate(tp('llsc', 1))} acq/s with 1 thread → {rate(tp('llsc', 8))} with 8 "
    f"— every try pulls the flag's cache line",
    f"One CPU, 8 threads: {rate(tp('llsc', 8, 'one-cpu'))} acq/s",
    "Waiters only read until the flag looks free, then try once (the LL/SC pair on Arm)",
], size=18)

r8 = first_run("llsc", 8, "one-cpu")
s = new_slide("LL/SC: Fairness and Limitations",
              f"Nothing orders the waiters: whichever store-conditional lands first wins. With "
              f"eight threads on one CPU the acquisition counts were severely unbalanced, though the scheduler plays a part. The lock also depends on "
              f"the hardware, which x86 does not have.")
text(s, 0.9, 1.7, 11.8, 5.4, [
    f"No queue: the first SC to land wins — fairness at 8 threads {fair('llsc', 8):.2f} (all CPUs), "
    f"{fair('llsc', 8, 'one-cpu'):.2f} (one CPU; run 1: the unluckiest thread got "
    f"{min(r8['counts']):,} of {r8['total']:,})",
    "A thread preempted while holding the lock makes every other thread spin out its time slice",
    "Limitations:",
    ("Starvation is possible; spins instead of sleeping", 1),
    ("Hardware-specific: not on x86; SC may fail spuriously", 1),
    ("Little may run between LL and SC, or the reservation is lost", 1),
    "Advantage over CAS: SC fails on any store, even one that writes back the same value "
    "(no ABA problem)",
], size=20)

# ===================================================================== MCS
mcs = [r for (l, _, _), rs in B.items() if l == "mcs" for r in rs]
section("MCS Lock",
        "Part three, the lock not covered in the lecture: the MCS queue lock.")

s = new_slide("MCS Lock: Idea",
              "The MCS lock, by Mellor-Crummey and Scott, won the Dijkstra Prize in 2006. "
              "Waiting threads form a queue, a linked list. Each thread spins on a flag in its "
              "own node, and the thread ahead of it clears that flag when it is done. Linux's "
              "spinlock today, the qspinlock, is built on this idea.")
text(s, 0.6, 1.7, 4.8, 5.4, [
    "Mellor-Crummey and Scott, 1991; Dijkstra Prize 2006",
    "Waiters form a queue: a linked list of nodes, tail points to the last",
    "Join: one atomic exchange on tail, then link in behind the previous node",
    "Each thread spins on its OWN node; the one ahead clears it on release",
    ("Linux's spinlock (qspinlock, since v4.2) is built on MCS", 0, ACCENT),
], size=18)
code(s, 5.6, 1.6, 7.3, 5.5, cut(SRC / "3_mcs.c", r"^struct node \{", r"^\}")
     + "\n\n" + cut(SRC / "3_mcs.c", r"^void lock_acquire"))

s = new_slide("MCS: Mutual Exclusion and Deadlock",
              f"The atomic exchange puts every thread in one queue, and only the head of that "
              f"queue may enter. A thread leaves by handing the lock to the node behind it, or by "
              f"resetting tail if nobody is waiting. We observed no lost update and no hang in "
              f"all {len(mcs)} runs.")
text(s, 0.6, 1.7, 4.8, 5.4, [
    "Mutual exclusion: the exchange gives every thread one place in one queue; only the "
    "head is inside, and it wakes exactly one node",
    "No deadlock: the holder always hands over — to the next node, or resets tail "
    "when nobody waits",
    "release() waits for a waiter that has swapped tail but not linked in yet",
    f"No violation observed in {len(mcs)} runs (1–8 threads, all CPUs and one CPU): "
    f"lost = 0, every run finished",
], size=17)
code(s, 5.6, 1.6, 7.3, 5.5, cut(SRC / "3_mcs.c", r"^void lock_release"))

s = new_slide("MCS: Performance",
              f"Every waiter spins on its own cache line, so a handover touches one line instead "
              f"of making every core fight over the same flag. On one CPU it collapses: the queue "
              f"hands the lock to a thread that is not running. The Linux kernel avoids that by "
              f"disabling preemption while a spinlock is held or awaited.")
chart(s, 0.6, 1.6, 7.0, 5.4, XL_CHART_TYPE.LINE_MARKERS, [str(t) for t in THREADS],
      [(f"MCS, all {NCPU} CPUs", [mps("mcs", t, "all-cpus") for t in THREADS]),
       ("MCS, one CPU", [mps("mcs", t, "one-cpu") for t in THREADS]),
       (f"TAS, all {NCPU} CPUs", [mps("tas", t, "all-cpus") for t in THREADS]),
       ("TAS, one CPU", [mps("tas", t, "one-cpu") for t in THREADS])],
      "threads")
text(s, 7.9, 1.7, 5.0, 5.3, [
    "Each waiter spins on its own cache line; a handover touches one line",
    f"All CPUs, 8 threads: {rate(tp('mcs', 8))} acq/s (TAS: {rate(tp('tas', 8))})",
    (f"One CPU, 8 threads: {rate(tp('mcs', 8, 'one-cpu'))} acq/s — the next in the "
     f"queue is not running", 0, ACCENT),
    "In the kernel, spinlocks run with preemption disabled, so a queued waiter is never "
    "descheduled",
], size=18)

m1 = first_run("mcs", 8, "one-cpu")
s = new_slide("MCS: Fairness and Limitations",
              "The queue makes MCS first-come, first-served: it was the fairest lock we measured "
              "at eight threads on all CPUs. That order is a property of the algorithm. On one "
              "CPU the acquisition counts were very unequal, but that is the scheduler: a queued "
              "thread that is not running cannot take its turn. MCS needs a node per waiting "
              "thread and an atomic exchange.")
text(s, 0.9, 1.7, 11.8, 5.4, [
    f"First-come, first-served through the queue: fairness {fair('mcs', 2):.2f} / "
    f"{fair('mcs', 4):.2f} / {fair('mcs', 8):.2f} at 2/4/8 threads (all CPUs) — the fairest "
    f"lock at 8 threads",
    "Limitations:",
    ("Needs an atomic exchange (and CAS for release) from the hardware", 1),
    ("One queue node per waiting thread, passed in or kept per thread", 1),
    ("A preempted waiter stalls everyone behind it — the kernel disables preemption", 1),
    (f"On one CPU here: {rate(tp('mcs', 8, 'one-cpu'))} acq/s at 8 threads, fairness "
     f"{fair('mcs', 8, 'one-cpu'):.2f} (run 1: {min(m1['counts']):,} vs {max(m1['counts']):,}) — "
     f"the scheduler, not the queue order", 1),
], size=20)

# ============================================================== comparison
def most(lock):
    """The most threads a lock was run with: Peterson stops at 2."""
    return 2 if lock == "peterson" else 8


def short(lock):
    return re.sub(r" [(].*", "", LABEL[lock])


cats = [str(t) for t in THREADS]

section("Comparison: Throughput",
        "Now the three locks side by side, with two locks from the lecture as references.")

s = new_slide(f"Throughput on All {NCPU} CPUs",
              f"With one thread nobody waits, so the simplest locks, test-and-set and the CAS "
              f"spin lock, are fastest. As threads are added every spin lock slows down, because the "
              f"threads fight over shared cache lines. pthread_mutex was the fastest lock at "
              f"eight threads.")
chart(s, 0.5, 1.5, 7.6, 5.6, XL_CHART_TYPE.LINE_MARKERS, cats,
      [(LABEL[l], [mps(l, t, "all-cpus") for t in THREADS]) for l in LOCKS], "threads")
lead1 = sorted(LOCKS, key=lambda l: -tp(l, 1))[:2]
text(s, 8.4, 1.7, 4.5, 5.3, [
    f"Medians of {REPS} runs. 1 thread, no contention: {short(lead1[0])} {rate(tp(lead1[0], 1))} and "
    f"{short(lead1[1])} {rate(tp(lead1[1], 1))} acq/s are fastest",
    "More threads: every spin lock slows down as threads fight over cache lines",
    f"8 threads: MCS {rate(tp('mcs', 8))}, CAS spin {rate(tp('llsc', 8))}, "
    f"TAS {rate(tp('tas', 8))} acq/s",
    (f"Fastest at 8 threads: pthread_mutex, {rate(tp('mutex', 8))} acq/s", 0, ACCENT),
], size=17)

s = new_slide("Throughput on One CPU",
              "Pinned to one CPU, the picture changes. Test-and-set and the CAS spin lock stay "
              "fast only because one thread keeps the lock for its whole time slice while the "
              "others get almost nothing. Peterson and MCS hand the lock to a thread that is "
              "not running, so they collapse. The mutex puts waiters to sleep and keeps its "
              "speed. These one-CPU numbers vary a lot between runs, as the variation slide shows.")
chart(s, 0.5, 1.5, 7.6, 5.6, XL_CHART_TYPE.LINE_MARKERS, cats,
      [(LABEL[l], [mps(l, t, "one-cpu") for t in THREADS]) for l in LOCKS], "threads")
text(s, 8.4, 1.7, 4.5, 5.3, [
    f"TAS / CAS spin stay fast ({rate(tp('tas', 8, 'one-cpu'))} / {rate(tp('llsc', 8, 'one-cpu'))} "
    f"at 8 threads), but one thread holds on while the others get almost nothing",
    f"Peterson and MCS hand the lock to a thread that is not running: "
    f"{rate(tp('peterson', 2, 'one-cpu'))} (Peterson, 2 thr), "
    f"{rate(tp('mcs', 8, 'one-cpu'))} (MCS, 8 thr)",
    (f"pthread_mutex sleeps instead of spinning: {rate(tp('mutex', 1, 'one-cpu'))} → "
     f"{rate(tp('mutex', 8, 'one-cpu'))} acq/s from 1 to 8 threads", 0, ACCENT),
], size=17)

section("Comparison: Summary",
        "Next, correctness and fairness for every lock, and then everything in one table.")

pet_bench = [r for (l, _, _), rs in B.items() if l == "peterson" for r in rs]
s = new_slide("Correctness: Mutual Exclusion and Deadlock",
              "For every lock we observed no violation of either correctness metric: no update "
              "was lost and every run finished. That is what the runs showed, not a proof. The "
              "one exception is the deliberate one, Peterson without its memory fence, which "
              "lost updates on every multi-CPU run.")
rows = [["Lock", "Runs", "Lost updates", "Runs finished", "Also tested"]]
extra = {
    "peterson": f"volatile, no fence: lost updates in {sum(r['lost'] > 0 for r in nofence)} / "
                f"{len(nofence)} runs",
    "llsc": f"x86 = CAS proxy; real ldaxr/stxr under qemu: lost = 0 in "
            f"{sum(r['lost'] == 0 for r in arm)} / {len(arm)} runs",
    "mcs": "", "tas": "", "mutex": "",
}
for l in LOCKS:
    rs = [r for (k, _, _), v in B.items() if k == l for r in v]
    if l == "peterson":
        rs += fenced + atomic
    rows.append([short(l), str(len(rs)), f"{sum(r['lost'] for r in rs):,}",
                 f"{len(rs)} / {len(rs)}", extra[l]])
table(s, 0.5, 1.7, 12.3, rows, [3.2, 0.9, 1.5, 1.6, 5.1], size=14, row_h=0.55)
text(s, 0.5, 1.7 + 0.55 * len(rows) + 0.4, 12.3, 1.5, [
    "Mutual exclusion: lost = total acquisitions − counter, summed over all runs of the correct builds",
    ("Peterson held on x86 only with the fence (or _Atomic seq_cst)", 0, ACCENT),
], size=16)

def counts_text(lock):
    """Fewest and most acquisitions in run 1 on one CPU, at the lock's most threads."""
    r = first_run(lock, most(lock), "one-cpu")
    return f"{min(r['counts']):,} vs {max(r['counts']):,}"


zero = sum(min(r["counts"]) == 0 for rs in B.values() for r in rs)
s = new_slide("Fairness",
              f"Fairness is the fewest acquisitions of any thread divided by the most. On all "
              f"CPUs the MCS queue was the fairest lock at eight threads. On one CPU every spin "
              f"lock showed severe acquisition imbalance during the measurement window. That "
              f"does not prove starvation: a thread that got one acquisition shows as 0.00 "
              f"here. Whether the algorithm is FIFO or has bounded waiting is a separate "
              f"question from what the scheduler did in one second.")
chart(s, 0.5, 1.5, 7.6, 5.6, XL_CHART_TYPE.COLUMN_CLUSTERED,
      [f"{short(l)} ({most(l)} thr)" for l in LOCKS],
      # rounded here so the bar labels match the text, which uses the same rounding
      [(f"all {NCPU} CPUs", [float(f"{fair(l, most(l)):.2f}") for l in LOCKS]),
       ("one CPU", [float(f"{fair(l, most(l), 'one-cpu'):.2f}") for l in LOCKS])],
      "lock (most threads it ran)", value_title="fairness (fewest ÷ most)", log=False,
      labels="0.00")
text(s, 8.4, 1.7, 4.5, 5.3, [
    f"All CPUs: MCS {fair('mcs', 8):.2f} is the fairest at 8 threads — its queue is "
    f"first-come, first-served",
    f"TAS {fair('tas', 8):.2f}, CAS spin {fair('llsc', 8):.2f}: whoever wins the race gets in",
    ("One CPU: severe imbalance for every spin lock — the scheduler decides who runs, "
     f"e.g. MCS {counts_text('mcs')}, TAS {counts_text('tas')} acquisitions (run 1, fewest vs most)",
     0, ACCENT),
    f"Only pthread_mutex stays balanced on one CPU: {fair('mutex', 8, 'one-cpu'):.2f}",
    f"Thread-runs with zero acquisitions: {zero} of "
    f"{sum(len(r['counts']) for rs in B.values() for r in rs)}",
], size=16)

def spread(lock, t, mode):
    xs = [r["throughput"] for r in B[(lock, t, mode)]]
    return max(xs) / min(xs)


s = new_slide("Run-to-Run Variation",
              f"A median hides how much runs differ. Here are the slowest and fastest of {REPS} "
              f"runs for each lock. The spread shows how far one number can be trusted, so we "
              f"read the results as the shape of the behaviour, not as values that would "
              f"repeat exactly.")
rows = [["Lock (threads)", "CPUs", "slowest run", "median", "fastest run", "max ÷ min"]]
for mode, name in (("all-cpus", f"all {NCPU}"), ("one-cpu", "one")):
    for l in LOCKS:
        t = most(l)
        rows.append([f"{short(l)} ({t})", name,
                     rate(min(r["throughput"] for r in B[(l, t, mode)])),
                     rate(tp(l, t, mode)),
                     rate(max(r["throughput"] for r in B[(l, t, mode)])),
                     f"{spread(l, t, mode):.1f}×"])
table(s, 0.5, 1.5, 12.3, rows, [4.3, 1.4, 1.6, 1.6, 1.6, 1.8], size=13, row_h=0.43)
text(s, 0.5, 1.5 + 0.43 * len(rows) + 0.15, 12.3, 0.8, [
    f"acq/s per run; {REPS} runs per setting, no confidence intervals",
], size=14)

s = new_slide("Summary",
              "This table puts the three locks and the two references on one page, each at the "
              "most threads it supports. Peterson is balanced but limited to two threads, the "
              "simple spin locks are unbalanced, and the MCS queue was balanced at every thread "
              "count we ran on all CPUs. The CAS spin lock is the x86 proxy for LL/SC.")
rows = [["Lock", "Needs", "Mutual excl.", "Deadlock-free", "Fairness",
         f"acq/s, all {NCPU} CPUs", "acq/s, one CPU"]]
needs = {"peterson": "loads/stores + seq_cst", "llsc": "CAS (LL/SC on Arm)",
         "mcs": "atomic exchange + CAS", "tas": "test-and-set", "mutex": "atomics + futex"}
for l in LOCKS:
    t = most(l)
    rows.append([f"{short(l)} ({t} thr)", needs[l], "yes", "yes",
                 f"{fair(l, t):.2f}", rate(tp(l, t)), rate(tp(l, t, "one-cpu"))])
table(s, 0.5, 1.6, 12.3, rows, [2.6, 2.6, 1.4, 1.5, 1.2, 1.6, 1.4], size=14, row_h=0.45)
text(s, 0.5, 1.6 + 0.45 * len(rows) + 0.4, 12.3, 1.5, [
    "Fairness = fewest ÷ most acquisitions per thread; medians of "
    f"{REPS} runs of {SECS} s; mutual exclusion and deadlock: no violation observed",
    "CAS spin lock = x86 stand-in for LL/SC: its speed is not the speed of LL/SC hardware",
    ("Peterson held on x86 only with the fence (or _Atomic seq_cst)", 0, ACCENT),
], size=16)

section("Conclusion",
        "Last part: what the measurements tell us, and where the material comes from.")
s = new_slide("Conclusion",
              "Each lock has an argument for mutual exclusion and deadlock freedom, and in our "
              "runs we observed no violation, though Peterson needs sequential consistency, a "
              "fence on today's x86. Peterson is balanced but limited to two threads; LL/SC is "
              "short and, as a CAS spin lock on x86, fast but unbalanced; the MCS queue is "
              "first-come, first-served for any number of threads and is what Linux builds its "
              "spinlock on. When threads can be preempted, sleeping locks such as pthread_mutex "
              "win. Our numbers compare locks with each other on one WSL2 machine.")
text(s, 0.9, 1.7, 11.8, 5.4, [
    f"Arguments plus no violation observed in {N_GOOD} runs — Peterson only with sequential consistency (a fence on x86)",
    "Peterson: software only, bounded waiting, but two threads only",
    "LL/SC: one short loop; our x86 numbers are a CAS proxy; no queue, so unbalanced",
    "MCS: a queue gives first-come, first-served for any number of threads; Linux builds on it",
    "Spinning on one CPU wastes time slices; sleeping locks (futex, pthread_mutex) avoid it",
    f"Limits: WSL2 vCPUs, tiny critical section, {REPS} × {SECS} s runs; some results vary widely",
], size=20)

s = new_slide("References",
              "These are the sources for the algorithms and the instructions.")
text(s, 0.9, 1.7, 11.8, 5.4, [
    "R. H. Arpaci-Dusseau, A. C. Arpaci-Dusseau, Operating Systems: Three Easy Pieces, ch. 28 “Locks”",
    "CPE 333 Lecture 8: Locks Mechanism",
    "G. L. Peterson, “Myths About the Mutual Exclusion Problem”, Information Processing Letters 12(3), 1981",
    "J. M. Mellor-Crummey, M. L. Scott, “Algorithms for Scalable Synchronization on "
    "Shared-Memory Multiprocessors”, ACM TOCS 9(1), 1991",
    "Linux kernel source: kernel/locking/qspinlock.c",
    "Arm Architecture Reference Manual for A-profile architecture: LDAXR, STXR",
    "Code and raw results: PS08/src, PS08/results",
], size=18)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides)} slides)")
